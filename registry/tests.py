import io

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .models import Beneficiary
from .services import find_duplicate_groups, import_csv
from .validators import validate_national_id, validate_phone


def make(national_id="123456789", **kw):
    defaults = dict(full_name="Test Person", district="Gaza City", household_size=4)
    defaults.update(kw)
    return Beneficiary.objects.create(national_id=national_id, **defaults)


class ValidatorTests(TestCase):
    def test_id_must_be_nine_digits(self):
        validate_national_id("123456789")
        for bad in ["12345678", "1234567890", "12345678a", ""]:
            with self.assertRaises(ValidationError):
                validate_national_id(bad)

    def test_phone_formats(self):
        for good in ["0597654321", "597654321", "+970597654321", "059 765-4321", ""]:
            validate_phone(good)
        with self.assertRaises(ValidationError):
            validate_phone("12345")


class DuplicateTests(TestCase):
    def test_same_phone_is_flagged(self):
        make("111111111", full_name="A One", phone="0561111111")
        make("222222222", full_name="B Two", phone="0561111111")
        groups = find_duplicate_groups()
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["reason"], "phone")

    def test_same_name_and_district_ignores_case_and_spacing(self):
        make("111111111", full_name="Ali  Hassan", district="Rafah")
        make("222222222", full_name="ali hassan", district="rafah")
        self.assertEqual(len(find_duplicate_groups()), 1)

    def test_same_name_other_district_is_not_duplicate(self):
        make("111111111", full_name="Ali Hassan", district="Rafah")
        make("222222222", full_name="Ali Hassan", district="Khan Younis")
        self.assertEqual(find_duplicate_groups(), [])


class ImportTests(TestCase):
    HEADER = "national_id,full_name,phone,district,household_size,status\n"

    def run_import(self, rows):
        return import_csv(io.StringIO(self.HEADER + rows))

    def test_valid_rows_are_created(self):
        r = self.run_import("123456789,Sara Ahmad,0599123456,Gaza City,5,active\n")
        self.assertEqual((r.created, len(r.errors)), (1, 0))
        self.assertEqual(Beneficiary.objects.get().status, "active")

    def test_bad_rows_are_reported_but_good_rows_are_kept(self):
        r = self.run_import(
            "123456789,Sara Ahmad,0599123456,Gaza City,5,active\n"
            "12345,Short Id,0599123456,Gaza City,3,active\n"
            "987654321,Bad Phone,123,Rafah,3,active\n"
            "555555555,,0599123456,Rafah,3,active\n"
        )
        self.assertEqual(r.created, 1)
        self.assertEqual([line for line, _ in r.errors], [3, 4, 5])

    def test_existing_ids_are_skipped_not_duplicated(self):
        make("123456789")
        r = self.run_import("123456789,Sara Ahmad,0599123456,Gaza City,5,active\n")
        self.assertEqual((r.created, r.skipped_existing), (0, 1))
        self.assertEqual(Beneficiary.objects.count(), 1)

    def test_invalid_status_is_rejected(self):
        r = self.run_import("123456789,Sara Ahmad,0599123456,Gaza City,5,banana\n")
        self.assertEqual(r.created, 0)
        self.assertEqual(len(r.errors), 1)


class ViewTests(TestCase):
    def setUp(self):
        make("111111111", full_name="Ali Hassan", household_size=6)
        make("222222222", full_name="Mona Saleh", household_size=3, status="active")

    def test_dashboard_numbers(self):
        r = self.client.get(reverse("registry:dashboard"))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.context["total"], 2)
        self.assertEqual(r.context["people_reached"], 9)

    def test_search_and_status_filter(self):
        r = self.client.get(reverse("registry:list"), {"q": "mona"})
        self.assertEqual([b.full_name for b in r.context["items"]], ["Mona Saleh"])
        r = self.client.get(reverse("registry:list"), {"status": "active"})
        self.assertEqual(len(r.context["items"]), 1)

    def test_create_rejects_invalid_id(self):
        r = self.client.post(
            reverse("registry:new"),
            {"national_id": "abc", "full_name": "X", "district": "Y", "household_size": 1, "status": "pending"},
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(Beneficiary.objects.count(), 2)

    def test_create_valid_redirects(self):
        r = self.client.post(
            reverse("registry:new"),
            {"national_id": "333333333", "full_name": "New One", "district": "Deir al-Balah", "household_size": 2, "status": "pending"},
        )
        self.assertRedirects(r, reverse("registry:list"))
        self.assertEqual(Beneficiary.objects.count(), 3)

    def test_excel_export_has_header_and_rows(self):
        r = self.client.get(reverse("registry:export"))
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws["A1"].value, "ID number")
